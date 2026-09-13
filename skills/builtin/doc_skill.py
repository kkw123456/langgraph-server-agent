"""文档处理技能：PDF 解析（pypdf）、Word 读取（python-docx）、Excel 读写（openpyxl）。"""
import os

from langchain_core.tools import tool

import workspace


def _safe_path(path: str) -> str:
    root = workspace.current_root()
    full = os.path.normpath(os.path.join(root, path))
    if full != root and not full.startswith(root + os.sep):
        raise ValueError("路径越界")
    return full


@tool
def parse_pdf(path: str, max_pages: int = 30) -> str:
    """提取工作目录下 PDF 文件的文本内容（前 max_pages 页，每页以 --- 第 N 页 --- 分隔）。

    Args:
        path: PDF 文件相对路径。
        max_pages: 最多解析的页数，默认 30。
    """
    try:
        from pypdf import PdfReader
    except ImportError:
        return "PDF 依赖未安装（pypdf），请在服务器执行: pip install pypdf"
    full = _safe_path(path)
    if not os.path.isfile(full):
        return f"文件不存在: {path}"
    try:
        reader = PdfReader(full)
        total = len(reader.pages)
        parts = []
        for i, page in enumerate(reader.pages[: max(1, max_pages)]):
            t = (page.extract_text() or "").strip()
            parts.append(f"--- 第 {i + 1} 页 ---\n{t}")
        out = "\n\n".join(parts)
        if total > max_pages:
            out += f"\n\n（共 {total} 页，仅解析前 {max_pages} 页）"
        return out[:8000] or "PDF 中未提取到文本（可能是扫描件，需要 OCR）"
    except Exception as e:
        return f"PDF 解析失败: {e}"


@tool
def read_docx(path: str) -> str:
    """读取工作目录下 Word（.docx）文件的段落与表格内容（前 8000 字符）。"""
    try:
        import docx
    except ImportError:
        return "Word 依赖未安装（python-docx），请在服务器执行: pip install python-docx"
    full = _safe_path(path)
    if not os.path.isfile(full):
        return f"文件不存在: {path}"
    try:
        d = docx.Document(full)
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for tb in d.tables:
            for row in tb.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)[:8000] or "（文档为空）"
    except Exception as e:
        return f"Word 读取失败: {e}"


def _md_table(rows, max_rows=30):
    if not rows:
        return "（空表）"
    head = "| " + " | ".join(str(c) for c in rows[0]) + " |"
    sep = "|" + "|".join(["---"] * len(rows[0])) + "|"
    body = "\n".join("| " + " | ".join("" if c is None else str(c) for c in r) + " |" for r in rows[1:max_rows])
    more = f"\n（仅显示前 {max_rows - 1} 行数据）" if len(rows) > max_rows else ""
    return f"{head}\n{sep}\n{body}{more}"


@tool
def excel_read(path: str, sheet: str = "", max_rows: int = 30) -> str:
    """读取工作目录下 Excel（.xlsx）文件为 Markdown 表格。

    Args:
        path: Excel 文件相对路径。
        sheet: 工作表名，默认第一个。
        max_rows: 最多返回的数据行数，默认 30。
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "Excel 依赖未安装（openpyxl），请在服务器执行: pip install openpyxl"
    full = _safe_path(path)
    if not os.path.isfile(full):
        return f"文件不存在: {path}"
    try:
        wb = load_workbook(full, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        names = "、".join(wb.sheetnames) if hasattr(wb, "sheetnames") else ""
        out = _md_table(rows, max(3, max_rows + 1))
        return f"工作表: {ws.title}（可选: {names}）\n\n{out}"
    except Exception as e:
        return f"Excel 读取失败: {e}"


@tool
def excel_write(data_json: str, path: str = "", sheet: str = "Sheet1") -> str:
    """把二维数据写入 Excel 文件（.xlsx），保存到工作目录。

    Args:
        data_json: 二维数组 JSON，第一行为表头，如 [["姓名","分数"],["张三",90],["李四",85]]。
        path: 目标文件相对路径，默认自动命名 data_<时间戳>.xlsx。
        sheet: 工作表名，默认 Sheet1。
    """
    try:
        import json as _json
        from openpyxl import Workbook
    except ImportError:
        return "Excel 依赖未安装（openpyxl）"
    try:
        data = _json.loads(data_json)
        if not isinstance(data, list) or not data or not isinstance(data[0], list):
            return "data_json 需为二维数组 JSON，如 [[\"列A\",\"列B\"],[1,2]]"
    except Exception as e:
        return f"数据解析失败: {e}"
    if not path:
        path = f"data_{int(__import__('time').time())}.xlsx"
    full = _safe_path(path)
    parent = os.path.dirname(full)
    if parent:
        os.makedirs(parent, exist_ok=True)
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = sheet or "Sheet1"
        for row in data[:5000]:
            ws.append(row)
        wb.save(full)
        return f"Excel 已写入: {path}（{len(data[:5000])} 行）"
    except Exception as e:
        return f"Excel 写入失败: {e}"


def get_skill():
    from skills.base import Skill
    return Skill(
        id="doc",
        name="文档处理",
        category="文档",
        description="PDF 文本解析、Word（docx）读取、Excel（xlsx）读取与写入，解析结果直接进入对话。",
        tools=[parse_pdf, read_docx, excel_read, excel_write],
        system_prompt=(
            "用户上传 PDF/Word/Excel 并要求分析时：先用 parse_pdf / read_docx / excel_read 读取内容，"
            "再基于内容回答；需要生成 Excel 报表时用 excel_write 并把返回路径告知用户。"
        ),
    )
